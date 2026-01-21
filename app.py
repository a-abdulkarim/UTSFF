import pandas as pd
import streamlit as st
import graphviz
from docx import Document
import io
import tempfile

# Framework steps and substeps
framework_steps = {
    "1. Data Collection": {
        "Identify relevant data sources": ["Search trends data", "Web traffic data", "Social media metrics", "External indicators", "Official sources"],
        "Define key variables and keywords": ["Relevant search keywords", "Specific event-based searches", "Behavioral indicators"],
        "Data acquisition methods": ["API extraction", "Web scraping", "Provided data"]
    },
    "2. Data Preprocessing": {
        "Cleaning and transformation": ["Handle missing values", "Remove duplicates", "Standardize formats"],
        "Feature engineering": ["Create lag variables", "Rolling averages", "Normalize indicators"],
        "Handling seasonality and trends": ["Add cyclical features", "Differencing if needed"]
    },
    "3. Feature Selection": {
        "Filtering variables": ["Correlation analysis", "Remove low-variance features"],
        "Dimensionality reduction": ["Recursive feature elimination", "PCA if beneficial"],
        "Selecting predictors": ["Assess feature importance", "Eliminate redundancy"]
    },
    "4. Model Selection": {
        "Choosing models": ["ARIMA, Exponential Smoothing", "XGBoost, Random Forest, Neural Networks"],
        "Hyperparameter tuning": ["Grid or random search", "Bayesian optimization"],
        "Complexity & interpretability": ["Simple vs complex comparison", "Computational efficiency"]
    },
    "5. Training & Testing": {
        "Splitting data": ["Train-test split", "Rolling forecast validation"],
        "Addressing anomalies": ["Handle outliers", "Adjust for unexpected events"],
        "Optimizing training": ["Early stopping", "Prevent overfitting"]
    },
    "6. Validation": {
        "Statistical validation": ["Evaluate MSE, RMSE, MAPE", "Check residuals"],
        "Robustness checks": ["Test scenarios", "Check consistency"]
    },
    "7. Evaluation": {
        "Accuracy assessment": ["Compare predictions vs actuals", "Monitor drift"],
        "Interpretation": ["Feature importance analysis", "Explainability tools like SHAP"]
    },
    "8. Deployment": {
        "Readiness": ["Automate data updates", "Monitoring dashboards"],
        "Continuous improvement": ["Retrain models", "Update features"]
    }
}

# Help text for each step and substep
help_text = {

    "1. Data Collection": {
        "Identify relevant data sources": "This step involves selecting relevant data sources such as search trends, web traffic, and social media metrics.",
        "Define key variables and keywords": "Identify the important variables and keywords that will be used for analysis and forecasting.",
        "Data acquisition methods": "Determine the method for acquiring data, such as through APIs, web scraping, or external data providers."
    },
    "2. Data Preprocessing": {
        "Cleaning and transformation": "Clean and transform the data to ensure it is ready for modeling, including handling missing values and removing duplicates.",
        "Feature engineering": "Create new features from the raw data, such as lag variables, rolling averages, or normalized indicators.",
        "Handling seasonality and trends": "Address seasonality and trends in the data by adding cyclical features or differencing the data if necessary."
    },
    "3. Feature Selection": {
        "Filtering variables": "Filter out irrelevant or redundant variables based on correlation or variance analysis.",
        "Dimensionality reduction": "Use techniques like recursive feature elimination or principal component analysis (PCA) to reduce the feature space.",
        "Selecting predictors": "Select the most important features for the model by assessing feature importance and eliminating redundancy."
    },
    "4. Model Selection": {
        "Choosing models": "Choose appropriate models for forecasting, such as ARIMA, Exponential Smoothing, or machine learning models like XGBoost and Random Forest.",
        "Hyperparameter tuning": "Tune the model's hyperparameters using techniques like grid search, random search, or Bayesian optimization.",
        "Complexity & interpretability": "Balance model complexity with interpretability, ensuring the model is both effective and understandable."
    },
    "5. Training & Testing": {
        "Splitting data": "Split the data into training and testing sets, or use rolling forecast validation for time series data.",
        "Addressing anomalies": "Identify and handle anomalies in the data, such as outliers or unexpected events that could affect model performance.",
        "Optimizing training": "Use techniques like early stopping to optimize training and prevent overfitting."
    },
    "6. Validation": {
        "Statistical validation": "Evaluate model performance using statistical metrics like MSE, RMSE, and MAPE, and check residuals for patterns.",
        "Robustness checks": "Test the robustness of the model by applying it to different scenarios and checking for consistency."
    },
    "7. Evaluation": {
        "Accuracy assessment": "Assess the accuracy of the model by comparing predictions with actual values and monitoring for drift over time.",
        "Interpretation": "Interpret the results of the model, including feature importance analysis and using tools like SHAP for explainability."
    },
    "8. Deployment": {
        "Readiness": "Ensure the model is ready for deployment, with automated data updates and monitoring dashboards in place.",
        "Continuous improvement": "Plan for continuous improvement of the model by retraining with new data and updating features as needed."
    }
}

# Session state initialization
if "checked_items" not in st.session_state:
    st.session_state.checked_items = {step: {sub: [False]*len(steps) for sub, steps in subcats.items()} for step, subcats in framework_steps.items()}

st.title("📊 Forecasting Framework Checklist")

completed_steps = {}

# Display each step with help text
for step, subcats in framework_steps.items():
    with st.expander(step):
        sub_completed = []
        for sub, steps_list in subcats.items():
            # Displaying the sub step with help icon inline
            st.markdown(f"**{sub}** <a href='#' title='{help_text[step][sub]}'> ❓</a>", unsafe_allow_html=True)
            
            # Regular checkboxes for predefined options
            checked = [st.checkbox(step_text, st.session_state.checked_items[step][sub][i]) for i, step_text in enumerate(steps_list)]
            st.session_state.checked_items[step][sub] = checked
            if any(checked):
                done_steps = [steps_list[i] for i, val in enumerate(checked) if val]
                sub_completed.append(f"{sub}: {', '.join(done_steps)}")
            
            # Add extra checkbox for custom input
            extra_option = st.checkbox("Other (Specify)", key=f"extra_{step}_{sub}")
            if extra_option:
                custom_input = st.text_input(f"Enter custom option for {sub}:", key=f"custom_{step}_{sub}")
                if custom_input:
                    sub_completed.append(f"{custom_input}")
                    
        if sub_completed:
            completed_steps[step] = sub_completed

def generate_diagram(completed):
    dot = graphviz.Digraph()
    dot.node("Start", shape="ellipse", style="filled", fillcolor="lightblue")
    prev = "Start"
    for step in framework_steps:
        label = step + ("\n" + "\n".join(completed.get(step, ["(Not done)"])) if step in completed else "\n(Not done)")
        color = "lightgray" if step in completed else "lightgreen"
        dot.node(step, label=label, shape="box", style="filled", fillcolor=color)
        dot.edge(prev, step)
        prev = step
    dot.node("End", shape="ellipse", style="filled", fillcolor="lightgreen")
    dot.edge(prev, "End")
    return dot

st.subheader("Framework Overview")
st.graphviz_chart(generate_diagram(completed_steps))

start_date = st.sidebar.date_input("Start Date")
end_date = st.sidebar.date_input("End Date")
frequency = st.sidebar.selectbox("Frequency", ["Daily", "Weekly", "Monthly", "Quarterly", "Yearly"], index=2)
horizon = st.sidebar.number_input("Forecast Horizon", min_value=1, value=6)

if start_date and end_date:
    periods = len(pd.date_range(start=start_date, end=end_date, freq=frequency[0]))
    train = int(0.8 * periods)
    test = periods - train
    st.sidebar.write(f"Training: {train}, Test: {test}")
    if test < horizon:
        st.sidebar.warning("Test set shorter than forecast horizon.")
if completed_steps:
    for step, items in completed_steps.items():
        st.sidebar.subheader(step)
        for item in items:
            st.sidebar.write(f"- {item}")
else:
    st.sidebar.write("No steps completed yet.")
if completed_steps:
    doc = Document()
    doc.add_heading('Completed Forecasting Steps', level=1)
    for step, sublist in completed_steps.items():
        doc.add_heading(step, level=2)
        for line in sublist:
            doc.add_paragraph(line, style='List Bullet')
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    st.download_button("Download Completed Steps (DOCX)", buffer, "completed_steps.docx")

if st.button("Export Diagram"):
    with tempfile.TemporaryDirectory() as tmp:
        diagram = generate_diagram(completed_steps)
        diagram.render(filename="diagram", directory=tmp, cleanup=True)
        with open(f"{tmp}/diagram.png", "rb") as f:
            st.download_button("Download Diagram (PNG)", f, "framework_diagram.png")
